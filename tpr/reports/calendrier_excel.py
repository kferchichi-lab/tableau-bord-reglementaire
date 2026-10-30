"""Export Excel du calendrier de contrôle."""

import io

import pandas as pd

from tpr.config import ORDRE_CI
from tpr.domain.calendrier import calculer_taux_realisation


def generer_calendrier_controle_excel(df_calendrier: pd.DataFrame, annee_reference: int) -> bytes:
    from openpyxl.chart import BarChart, DoughnutChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.marker import DataPoint
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    df_export = df_calendrier[[c for c in df_calendrier.columns if not c.startswith("_")]]
    taux = calculer_taux_realisation(df_calendrier)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_export.to_excel(writer, index=False, sheet_name="Calendrier", startrow=1)
        ws = writer.sheets["Calendrier"]

        nb_cols = len(df_export.columns)
        ws.cell(row=1, column=1, value=f"Calendrier de contrôle réglementaire — Année {annee_reference}")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nb_cols)
        ws.cell(row=1, column=1).font = Font(size=14, bold=True, color="1E3A8A")
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

        header_row = 2
        header_fill = PatternFill("solid", fgColor="1E3A8A")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="CBD5E1")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col_idx in range(1, nb_cols + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        for row_idx in range(header_row + 1, header_row + 1 + len(df_export)):
            for col_idx in range(1, nb_cols + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        nb_installations = len(ORDRE_CI)
        for i in range(0, len(df_export), nb_installations):
            r1 = header_row + 1 + i
            r2 = header_row + nb_installations + i
            ws.merge_cells(start_row=r1, start_column=1, end_row=r2, end_column=1)
            ws.cell(row=r1, column=1).font = Font(bold=True, color="1E3A8A")

        largeurs = [10, 6, 26, 12, 10, 26, 18, 26]
        for i, largeur in enumerate(largeurs, start=1):
            ws.column_dimensions[get_column_letter(i)].width = largeur

        # ================= FEUILLE "Taux de réalisation" (graphiques) =================
        ws_g = writer.book.create_sheet("Taux de réalisation")
        ws_g["A1"] = f"Taux de réalisation des contrôles — {annee_reference}"
        ws_g.merge_cells("A1:D1")
        ws_g["A1"].font = Font(size=14, bold=True, color="1E3A8A")
        ws_g["A1"].alignment = Alignment(horizontal="center")

        # --- Données pour la barre horizontale (par site) ---
        ws_g["A3"] = "Site"; ws_g["B3"] = "Taux (%)"
        ws_g["A3"].font = header_font; ws_g["A3"].fill = header_fill
        ws_g["B3"].font = header_font; ws_g["B3"].fill = header_fill
        ws_g["A4"] = "MEG"; ws_g["B4"] = taux.get("MEG", 0.0)
        ws_g["A5"] = "SGB"; ws_g["B5"] = taux.get("SGB", 0.0)

        bar = BarChart()
        bar.type = "bar"  # barres horizontales
        bar.title = f"Taux de réalisation par site — {annee_reference}"
        bar.y_axis.title = None
        bar.x_axis.title = "Taux (%)"
        bar.x_axis.scaling.min = 0
        bar.x_axis.scaling.max = 100
        data_bar = Reference(ws_g, min_col=2, min_row=3, max_row=5)
        cats_bar = Reference(ws_g, min_col=1, min_row=4, max_row=5)
        bar.add_data(data_bar, titles_from_data=True)
        bar.set_categories(cats_bar)
        serie_bar = bar.series[0]
        pt_meg = DataPoint(idx=0); pt_meg.graphicalProperties.solidFill = "2563EB"
        pt_sgb = DataPoint(idx=1); pt_sgb.graphicalProperties.solidFill = "059669"
        serie_bar.data_points = [pt_meg, pt_sgb]
        serie_bar.dLbls = DataLabelList()
        serie_bar.dLbls.showVal = True
        bar.height, bar.width = 8.5, 16
        bar.legend = None
        ws_g.add_chart(bar, "D3")

        # --- Données pour l'anneau (taux global) ---
        ws_g["A8"] = "Répartition"; ws_g["B8"] = "Valeur"
        ws_g["A8"].font = header_font; ws_g["A8"].fill = header_fill
        ws_g["B8"].font = header_font; ws_g["B8"].fill = header_fill
        ws_g["A9"] = "Réalisé"; ws_g["B9"] = taux.get("Global", 0.0)
        ws_g["A10"] = "Restant"; ws_g["B10"] = round(100 - taux.get("Global", 0.0), 1)

        doughnut = DoughnutChart()
        doughnut.title = f"Taux global {annee_reference} : {taux.get('Global', 0.0)}%"
        data_don = Reference(ws_g, min_col=2, min_row=9, max_row=10)
        cats_don = Reference(ws_g, min_col=1, min_row=9, max_row=10)
        doughnut.add_data(data_don, titles_from_data=False)
        doughnut.set_categories(cats_don)
        serie_don = doughnut.series[0]
        pt_real = DataPoint(idx=0); pt_real.graphicalProperties.solidFill = "1E3A8A"
        pt_reste = DataPoint(idx=1); pt_reste.graphicalProperties.solidFill = "E2E8F0"
        serie_don.data_points = [pt_real, pt_reste]
        doughnut.height, doughnut.width = 8.5, 10
        ws_g.add_chart(doughnut, "D21")

        ws_g.column_dimensions["A"].width = 14
        ws_g.column_dimensions["B"].width = 12

    output.seek(0)
    return output.getvalue()
